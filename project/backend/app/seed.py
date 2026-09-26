"""Idempotent data loader. Run with: python -m app.seed

Reads pre-built data assets from DATA_DIR (see data/README.md for the expected
file layout and schemas) and loads them into the database. Safe to re-run —
existing rows are matched by their natural key and skipped.
"""

import asyncio
import json
import os
from pathlib import Path

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import hash_password
from app.database import AsyncSessionLocal, engine
from app.models.category import Category, Subcategory
from app.models.complaint import Complaint
from app.models.department import Department
from app.models.enums import ComplaintChannel, ComplaintStatus, LoyaltyTier, UserRole
from app.models.escalation_rule import EscalationRule
from app.models.knowledge_base import KnowledgeBaseDocument
from app.models.resolution_rule import ResolutionRule
from app.models.user import User
from app.utils.complaint_id import next_complaint_id

DATA_DIR = Path(os.environ.get("SUPPORTNOVA_DATA_DIR", Path(__file__).resolve().parent.parent / "data"))

DEMO_USERS = [
    # (full_name, email, password, role, department_name)
    ("Admin User", "admin@travelnova.com", "admin123", UserRole.ADMIN, None),
    ("Manager User", "manager@travelnova.com", "manager123", UserRole.MANAGER, None),
    ("Reviewer User", "reviewer@travelnova.com", "reviewer123", UserRole.REVIEWER, None),
    ("Flights Agent", "agent.flights@travelnova.com", "agent123", UserRole.AGENT, "Flight Operations"),
    ("Hotels Agent", "agent.hotels@travelnova.com", "agent123", UserRole.AGENT, "Hotel Services"),
    ("Billing Agent", "agent.billing@travelnova.com", "agent123", UserRole.AGENT, "Billing & Payments"),
    ("Demo Customer", "customer@example.com", "customer123", UserRole.CUSTOMER, None),
]

LOYALTY_TIER_MAP = {"silver": LoyaltyTier.SILVER, "gold": LoyaltyTier.GOLD, "platinum": LoyaltyTier.PLATINUM}


def _load_json(path: Path) -> list | None:
    if not path.exists():
        print(f"  skip (not found): {path}")
        return None
    with path.open(encoding="utf-8") as f:
        return json.load(f)


async def seed_categories(db: AsyncSession) -> int:
    print("Loading categories...")
    data = _load_json(DATA_DIR / "config" / "categories.json")
    if not data:
        return 0

    count = 0
    for cat_data in data:
        category = await db.scalar(select(Category).where(Category.code == cat_data["code"]))
        if category is None:
            category = Category(
                code=cat_data["code"],
                name=cat_data["name"],
                description=cat_data.get("description"),
            )
            db.add(category)
            await db.flush()
            count += 1

        for sub_data in cat_data.get("subcategories", []):
            subcategory = await db.scalar(select(Subcategory).where(Subcategory.code == sub_data["code"]))
            if subcategory is None:
                db.add(
                    Subcategory(
                        category_id=category.id,
                        code=sub_data["code"],
                        name=sub_data["name"],
                        description=sub_data.get("description"),
                    )
                )
                count += 1

    await db.commit()
    return count


async def seed_departments(db: AsyncSession) -> int:
    print("Loading departments...")
    data = _load_json(DATA_DIR / "config" / "departments.json")
    if not data:
        return 0

    count = 0
    for dept_data in data:
        existing = await db.scalar(select(Department).where(Department.code == dept_data["code"]))
        if existing is None:
            db.add(
                Department(
                    code=dept_data["code"],
                    name=dept_data["name"],
                    description=dept_data.get("description"),
                )
            )
            count += 1

    await db.commit()
    return count


async def seed_resolution_rules(db: AsyncSession) -> int:
    print("Loading resolution rules...")
    data = _load_json(DATA_DIR / "rules" / "complaint_resolution_rule_matrix.json")
    if not data:
        return 0

    count = 0
    for rule_data in data:
        existing = await db.scalar(select(ResolutionRule).where(ResolutionRule.rule_id == rule_data["rule_id"]))
        if existing is None:
            db.add(
                ResolutionRule(
                    rule_id=rule_data["rule_id"],
                    category=rule_data["category"],
                    subcategory=rule_data["subcategory"],
                    conditions=rule_data.get("conditions"),
                    department=rule_data["department"],
                    urgency=rule_data["urgency"],
                    priority=rule_data["priority"],
                    policy_id=rule_data.get("policy_id"),
                    escalation_required=rule_data.get("escalation_required", False),
                    escalation_level=rule_data.get("escalation_level", 0),
                    required_actions=rule_data.get("required_actions"),
                    prohibited_actions=rule_data.get("prohibited_actions"),
                    follow_up=rule_data.get("follow_up"),
                    compensation_eligible=rule_data.get("compensation_eligible", False),
                    refund_eligible=rule_data.get("refund_eligible", False),
                )
            )
            count += 1

    await db.commit()
    return count


async def seed_escalation_rules(db: AsyncSession) -> int:
    print("Loading escalation rules...")
    data = _load_json(DATA_DIR / "rules" / "escalation_rules.json")
    if not data:
        return 0

    count = 0
    for rule_data in data:
        existing = await db.scalar(select(EscalationRule).where(EscalationRule.rule_id == rule_data["rule_id"]))
        if existing is None:
            db.add(
                EscalationRule(
                    rule_id=rule_data["rule_id"],
                    trigger_condition=rule_data["trigger_condition"],
                    level=rule_data["level"],
                    level_name=rule_data["level_name"],
                    response_time=rule_data["response_time"],
                    is_mandatory=rule_data.get("is_mandatory", False),
                )
            )
            count += 1

    await db.commit()
    return count


async def seed_knowledge_base(db: AsyncSession) -> int:
    print("Indexing knowledge base documents...")
    policies_dir = DATA_DIR / "policies"
    if not policies_dir.exists():
        print(f"  skip (not found): {policies_dir}")
        return 0

    count = 0
    for docx_path in sorted(policies_dir.glob("*.docx")):
        document_id = docx_path.stem
        existing = await db.scalar(
            select(KnowledgeBaseDocument).where(KnowledgeBaseDocument.document_id == document_id)
        )
        if existing is not None:
            continue

        doc = DocxDocument(str(docx_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        title = paragraphs[0] if paragraphs else document_id
        content_text = "\n".join(paragraphs)

        db.add(
            KnowledgeBaseDocument(
                document_id=document_id,
                title=title,
                file_path=str(docx_path),
                content_text=content_text,
            )
        )
        count += 1

    await db.commit()
    return count


async def seed_users(db: AsyncSession) -> int:
    print("Creating demo users...")
    count = 0
    for full_name, email, password, role, department_name in DEMO_USERS:
        existing = await db.scalar(select(User).where(User.email == email))
        if existing is not None:
            continue

        department_id = None
        if department_name:
            department = await db.scalar(select(Department).where(Department.name == department_name))
            department_id = department.id if department else None

        db.add(
            User(
                email=email,
                password_hash=hash_password(password),
                full_name=full_name,
                role=role,
                department_id=department_id,
                loyalty_tier=LoyaltyTier.SILVER if role == UserRole.CUSTOMER else None,
            )
        )
        count += 1

    await db.commit()
    return count


async def _get_or_create_customer(db: AsyncSession, name: str, email: str) -> User:
    customer = await db.scalar(select(User).where(User.email == email))
    if customer is None:
        customer = User(
            email=email,
            password_hash=hash_password("customer123"),
            full_name=name,
            role=UserRole.CUSTOMER,
        )
        db.add(customer)
        await db.flush()
    return customer


async def seed_complaints(db: AsyncSession) -> int:
    print("Loading complaints...")
    data = _load_json(DATA_DIR / "complaints" / "complaints_dataset.json")
    if not data:
        return 0

    count = 0
    for item in data:
        customer = await _get_or_create_customer(db, item["customer_name"], item["email"])

        existing = await db.scalar(
            select(Complaint).where(
                Complaint.customer_id == customer.id,
                Complaint.title == item["title"],
                Complaint.description == item["description"],
            )
        )
        if existing is not None:
            continue

        loyalty_tier = LOYALTY_TIER_MAP.get((item.get("loyalty_tier") or "").lower())
        if loyalty_tier and customer.loyalty_tier != loyalty_tier:
            customer.loyalty_tier = loyalty_tier

        db.add(
            Complaint(
                complaint_id=await next_complaint_id(db),
                customer_id=customer.id,
                title=item["title"],
                description=item["description"],
                channel=ComplaintChannel.WEB_FORM,
                product_type=item["product_type"],
                booking_reference=item.get("booking_reference"),
                customer_selected_category=item.get("customer_selected_category"),
                status=ComplaintStatus.SUBMITTED,
            )
        )
        count += 1

        if count % 50 == 0:
            await db.commit()

    await db.commit()
    return count


async def print_summary(db: AsyncSession) -> None:
    counts = {
        "categories": len((await db.scalars(select(Category))).all()),
        "subcategories": len((await db.scalars(select(Subcategory))).all()),
        "departments": len((await db.scalars(select(Department))).all()),
        "resolution_rules": len((await db.scalars(select(ResolutionRule))).all()),
        "escalation_rules": len((await db.scalars(select(EscalationRule))).all()),
        "knowledge_base_documents": len((await db.scalars(select(KnowledgeBaseDocument))).all()),
        "users": len((await db.scalars(select(User))).all()),
        "complaints": len((await db.scalars(select(Complaint))).all()),
    }
    print("\n=== Seed summary ===")
    for key, value in counts.items():
        print(f"  {key}: {value}")


async def main() -> None:
    engine.echo = False

    async with AsyncSessionLocal() as db:
        new_categories = await seed_categories(db)
        new_departments = await seed_departments(db)
        new_rules = await seed_resolution_rules(db)
        new_escalation_rules = await seed_escalation_rules(db)
        new_docs = await seed_knowledge_base(db)
        new_users = await seed_users(db)
        new_complaints = await seed_complaints(db)

        print(
            f"\nInserted: {new_categories} categories/subcategories rows, {new_departments} departments, "
            f"{new_rules} resolution rules, {new_escalation_rules} escalation rules, {new_docs} kb docs, "
            f"{new_users} users, {new_complaints} complaints."
        )

        await print_summary(db)


if __name__ == "__main__":
    asyncio.run(main())
